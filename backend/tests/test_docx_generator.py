import io
import sys
import unittest

from docx import Document

from pathlib import Path

BACKEND_ROOT = Path(__file__).resolve().parents[1]
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))

from generators.docx_generator import generate_docx  # noqa: E402


class ClientTemplateDocxTests(unittest.TestCase):
    def test_client_template_docx_replaces_placeholders_and_appends_content(self):
        template = Document()
        template.add_paragraph("Project: {{project_name}}")
        template.add_paragraph("Client: {{client_name}}")

        buffer = io.BytesIO()
        template.save(buffer)

        content = {
          "metadata": {
            "project_name": "Atlas Revamp",
            "client_name": "Northwind Health",
          },
          "executive_summary": "Executive summary content.",
          "project_overview": "Overview content.",
          "scope_of_work": {
            "in_scope": ["Portal refresh"],
            "out_of_scope": ["Legacy migration"],
          },
          "deliverables": [{"id": "D-001", "name": "Delivery Pack", "description": "Packaged outputs"}],
          "risks": [{"id": "R-001", "description": "Key dependency risk"}],
          "assumptions": [{"id": "A-001", "description": "Stakeholder availability"}],
          "dependencies": [{"id": "DEP-001", "description": "API credentials"}],
          "responsibilities": {"Client": ["Approve milestones"], "Invuric": ["Deliver weekly status"]},
          "timeline": {"phases": [{"phase": "Discovery", "weeks": "1-2", "deliverables": "Findings"}]},
          "roles": [{"role": "BA", "description": "Leads discovery", "location": "Remote", "hours": "40"}],
          "change_management": "Formal review required.",
          "acceptance_criteria": ["Signed off by sponsor"],
          "validity": "30 days",
        }

        output = io.BytesIO()
        generate_docx(
            content,
            output,
            doc_type="SOW",
            template_mode="client",
            template_bytes=buffer.getvalue(),
            template_kind="docx",
        )

        generated = Document(io.BytesIO(output.getvalue()))
        text = "\n".join(paragraph.text for paragraph in generated.paragraphs if paragraph.text.strip())

        self.assertIn("Project: Atlas Revamp", text)
        self.assertIn("Client: Northwind Health", text)
        self.assertIn("SOW Generated Content", text)
        self.assertIn("Executive Summary", text)
        self.assertIn("Executive summary content.", text)

    def test_pdf_template_reference_still_generates_professional_document(self):
        content = {
          "metadata": {
            "project_name": "Atlas Revamp",
            "client_name": "Northwind Health",
          },
          "product_overview": "A modernized operating model for the client portal.",
          "problem_statement": "The current intake process is fragmented and slow.",
          "goals": [{"goal": "Reduce intake time", "metric": "Cycle time", "target": "< 10 min", "timeframe": "90 days"}],
          "user_personas": [{"name": "Priya", "role": "Operations Lead", "age_range": "30-40", "context": "Runs daily intake", "goals": "Faster processing", "pain_points": "Manual handoffs", "tech_proficiency": "Intermediate", "success_definition": "Same-day completion"}],
          "feature_requirements": [{"id": "F-001", "feature": "Guided intake", "description": "Step-by-step intake workflow", "priority": "High", "user_story": "As Priya, I want guided intake so that I can reduce errors.", "acceptance_criteria": "Scenario 1 - Given valid data, When submitted, Then the intake is saved. Scenario 2 - Given missing fields, When submitted, Then validation errors are shown.", "dependencies": "CRM API"}],
          "non_functional_requirements": [{"category": "Performance", "requirement": "Fast load time", "metric": "P95 < 2s"}],
          "user_flows": [{"flow_name": "Create request", "actor": "Priya", "steps": ["1. Open intake page", "2. Enter request details"], "success_outcome": "Request submitted"}],
          "edge_cases": [{"scenario": "Timeout", "trigger": "API timeout", "expected_system_behavior": "Retry gracefully", "user_communication": "Please try again"}],
          "risks": [{"id": "R-001", "risk": "Late API access", "likelihood": "Medium", "impact": "High", "mitigation": "Start onboarding in week 1"}],
          "technical_constraints": ["Must integrate with the existing CRM"],
          "out_of_scope": ["Native mobile apps"],
          "testing_strategy": {"approach": "Risk-based", "test_types": ["unit", "integration"], "coverage_areas": ["intake"], "vague_input_handling": "Prompt for missing details", "conflicting_requirements_handling": "Escalate to PO"},
          "assumptions": ["Stakeholders are available weekly"],
          "timeline": "12 weeks",
          "open_questions": ["Should intake support offline mode?"],
        }

        output = io.BytesIO()
        generate_docx(
            content,
            output,
            doc_type="PRD",
            template_mode="client",
            template_bytes=b"%PDF-1.4 template reference",
            template_kind="pdf",
        )

        generated = Document(io.BytesIO(output.getvalue()))
        text = "\n".join(paragraph.text for paragraph in generated.paragraphs if paragraph.text.strip())

        self.assertIn("PRD Generated Content", text)
        self.assertIn("Product Overview", text)
        self.assertIn("A modernized operating model for the client portal.", text)


if __name__ == "__main__":
    unittest.main()
