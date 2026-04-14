"""
Invuric BA Agent — Document Generator
- SOW: fills the actual Invuric SOW Word template (exact branding preserved)
- PRD / FRD: built from scratch matching the SOW template's exact style
"""

import copy, io, os, re, shutil
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from datetime import datetime

FONT = "Neue Haas Grotesk Text Pro"
TEMPLATES_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "templates")
INVURIC_LOGO_PATH = os.path.join(TEMPLATES_DIR, "invuric-logo.png")

HEX_BLACK  = "000000"
HEX_WHITE  = "FFFFFF"
HEX_LIGHT  = "F2F2F2"
HEX_ACCENT = "1A1A2E"   # dark navy for sub-headers in scratch docs
HEX_NAVY   = "1E3A5F"   # professional navy for cover page header band
HEX_SLATE  = "F0F4F8"   # light cool-gray for cover page footer band
GANTT_CHAR = "\u2588\u2588\u2588\u2588\u2588\u2588"

GANTT_COLS = [(1,2),(3,4),(5,6),(7,8),(9,10),(11,12),(13,16)]


# ── Low-level XML helpers ─────────────────────────────────────────────────────

def _safe_str(val) -> str:
    """Convert any value to string safely — None becomes empty string."""
    if val is None:
        return ""
    if isinstance(val, list):
        return "; ".join(_safe_str(v) for v in val)
    return str(val)


def _set_cell_text(cell, text):
    """Set cell text, preserving existing run formatting. Never writes 'None'."""
    text = _safe_str(text)
    for para in cell.paragraphs:
        if para.runs:
            para.runs[0].text = text
            for r in para.runs[1:]:
                r.text = ""
            return
    if cell.paragraphs:
        cell.paragraphs[0].add_run(text)
    else:
        cell.add_paragraph(text)


def _set_para_text(para, text):
    """Set paragraph text, preserving run formatting of first run."""
    text = _safe_str(text)
    if para.runs:
        para.runs[0].text = text
        for r in para.runs[1:]:
            r.text = ""
    else:
        para.add_run(text)


def _set_el_run_text(el, text):
    """Set text in a paragraph XML element (for cloned nodes)."""
    text = _safe_str(text)
    runs = el.findall(".//" + qn("w:r"))
    if runs:
        first = runs[0]
        for t in first.findall(qn("w:t")):
            first.remove(t)
        t_el = OxmlElement("w:t")
        t_el.text = text
        first.append(t_el)
        for r in runs[1:]:
            for t in r.findall(qn("w:t")):
                r.remove(t)
    else:
        r_el = OxmlElement("w:r")
        t_el = OxmlElement("w:t")
        t_el.text = text
        r_el.append(t_el)
        el.append(r_el)


def _replace_paras_with_items(paras: list, items: list):
    """Replace template list-paragraphs with AI items (handles different counts)."""
    if not paras:
        return
    items = [_safe_str(i) for i in items if i]

    ref_el = paras[0]._element
    last_el = paras[-1]._element

    for i, p in enumerate(paras):
        if i < len(items):
            _set_para_text(p, items[i])
        else:
            _set_para_text(p, "")

    if len(items) > len(paras):
        insert_after = last_el
        for extra in items[len(paras):]:
            new_el = copy.deepcopy(ref_el)
            _set_el_run_text(new_el, extra)
            insert_after.addnext(new_el)
            insert_after = new_el


def _fill_list_group(doc, pattern: str, items: list):
    """Find paragraphs matching pattern and replace with items."""
    paras = [p for p in doc.paragraphs if re.match(pattern, p.text.strip())]
    _replace_paras_with_items(paras, items)


def _shade(cell, hex_fill: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:shd")):
        tcPr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _borders(table, color=HEX_BLACK):
    for row in table.rows:
        for cell in row.cells:
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            for old in tcPr.findall(qn("w:tcBorders")):
                tcPr.remove(old)
            tcB = OxmlElement("w:tcBorders")
            for side in ("top", "left", "bottom", "right"):
                b = OxmlElement(f"w:{side}")
                b.set(qn("w:val"), "single")
                b.set(qn("w:sz"), "4")
                b.set(qn("w:space"), "0")
                b.set(qn("w:color"), color)
                tcB.append(b)
            tcPr.append(tcB)


def _hex(h: str) -> RGBColor:
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _style_run(run, size=11, bold=False, color_hex=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    if color_hex:
        run.font.color.rgb = _hex(color_hex)


def _clear_paragraph(paragraph):
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)


def _set_paragraph_text(paragraph, text: str, size=11, bold=False, color_hex=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    _clear_paragraph(paragraph)
    paragraph.alignment = align
    run = paragraph.add_run(_safe_str(text))
    _style_run(run, size=size, bold=bold, color_hex=color_hex)
    return paragraph


def _add_invuric_logo(paragraph, width=Inches(3.2)):
    _clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if os.path.exists(INVURIC_LOGO_PATH):
        paragraph.add_run().add_picture(INVURIC_LOGO_PATH, width=width)
    else:
        fallback = paragraph.add_run("Invuric")
        _style_run(fallback, size=28, bold=True)


def _header_row(table, headers, bg=HEX_BLACK, fg=HEX_WHITE, size=10):
    row = table.rows[0]
    for i, h in enumerate(headers):
        if i < len(row.cells):
            _shade(row.cells[i], bg)
            _set_cell_text(row.cells[i], h)
            for para in row.cells[i].paragraphs:
                for run in para.runs:
                    _style_run(run, size=size, bold=True, color_hex=fg)


def _set_col_widths(table, widths_in):
    for i, w in enumerate(widths_in):
        if i < len(table.columns):
            for cell in table.columns[i].cells:
                cell.width = Inches(w)


# ── Week helpers ──────────────────────────────────────────────────────────────

def _parse_week_nums(weeks_str: str, idx: int, total: int):
    if weeks_str:
        nums = re.findall(r"\d+", str(weeks_str))
        if len(nums) >= 2:
            return int(nums[0]), int(nums[1])
        if len(nums) == 1:
            w = int(nums[0])
            return w, w + 1
    chunk = 16 / max(total, 1)
    return int(idx * chunk) + 1, int((idx + 1) * chunk)


def _weeks_overlap(s1, e1, s2, e2) -> bool:
    return s1 <= e2 and s2 <= e1


# ═══════════════════════════════════════════════════════════════════════════════
#  SOW — template-copy approach
# ═══════════════════════════════════════════════════════════════════════════════

def _sow_from_template(content: dict, output_path: str):
    template_path = os.path.join(TEMPLATES_DIR, "invuric_sow.docx")
    shutil.copy2(template_path, output_path)
    doc = Document(output_path)
    metadata = content.get("metadata", {})

    _sow_cover(doc, metadata)
    _sow_metadata_table(doc, metadata)
    _sow_overview(doc, content)
    _sow_scope(doc, content)
    _sow_raid(doc, content)
    _sow_responsibilities(doc, content)
    _sow_gantt(doc, content)
    _sow_roles(doc, content)
    _sow_deliverables(doc, content)
    _sow_signature(doc, metadata)

    doc.save(output_path)


def _sow_cover(doc, metadata):
    tbl = doc.tables[0]
    cover_cell = tbl.rows[0].cells[2]
    client = metadata.get("client_name", "")
    today  = datetime.utcnow().strftime("%d %B %Y")
    paras = cover_cell.paragraphs

    for para in paras:
        _clear_paragraph(para)

    # Logo
    logo_para = paras[0]
    _add_invuric_logo(logo_para, width=Inches(3.2))
    logo_para.paragraph_format.space_after = Pt(24)

    # Document type — navy, large, bold
    title_para = paras[1] if len(paras) > 1 else cover_cell.add_paragraph()
    _set_paragraph_text(
        title_para,
        "STATEMENT OF WORK",
        size=24,
        bold=True,
        color_hex=HEX_NAVY,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    title_para.paragraph_format.space_after = Pt(2)

    # Subtitle tag line
    sub_para = paras[2] if len(paras) > 2 else cover_cell.add_paragraph()
    _set_paragraph_text(
        sub_para,
        "SOW",
        size=10,
        color_hex="94A3B8",
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    sub_para.paragraph_format.space_after = Pt(20)

    # "Prepared for" label — small, muted
    prepared_for_label = paras[3] if len(paras) > 3 else cover_cell.add_paragraph()
    _set_paragraph_text(
        prepared_for_label,
        "PREPARED FOR",
        size=9,
        color_hex="94A3B8",
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    prepared_for_label.paragraph_format.space_after = Pt(4)

    # Client name — navy, prominent
    client_para = paras[4] if len(paras) > 4 else cover_cell.add_paragraph()
    _set_paragraph_text(
        client_para,
        (client or "Client Name").upper(),
        size=18,
        bold=True,
        color_hex=HEX_NAVY,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    client_para.paragraph_format.space_after = Pt(20)

    # Footer meta — muted slate
    prepared_by_para = paras[5] if len(paras) > 5 else cover_cell.add_paragraph()
    _set_paragraph_text(
        prepared_by_para,
        "Invuric Technologies Pty. Ltd.",
        size=9,
        bold=True,
        color_hex="475569",
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    prepared_by_para.paragraph_format.space_after = Pt(3)

    date_para = paras[6] if len(paras) > 6 else cover_cell.add_paragraph()
    _set_paragraph_text(
        date_para,
        today,
        size=9,
        color_hex="94A3B8",
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )


def _sow_metadata_table(doc, metadata):
    tbl = doc.tables[1]
    today = datetime.utcnow().strftime("%d/%m/%Y")
    _set_cell_text(tbl.rows[0].cells[1], metadata.get("project_id", ""))
    _set_cell_text(tbl.rows[0].cells[3], metadata.get("project_name", ""))
    _set_cell_text(tbl.rows[1].cells[1], metadata.get("author", ""))
    _set_cell_text(tbl.rows[1].cells[3], metadata.get("requestor", ""))
    _set_cell_text(tbl.rows[2].cells[1], today)
    _set_cell_text(tbl.rows[2].cells[3], today)
    if len(tbl.rows) > 3:
        _set_cell_text(tbl.rows[3].cells[1], metadata.get("start_date", ""))
        _set_cell_text(tbl.rows[3].cells[3], metadata.get("end_date", ""))


def _sow_overview(doc, content):
    overview = content.get("project_overview", content.get("executive_summary", ""))
    for p in doc.paragraphs:
        if "<Overview Description>" in p.text:
            _set_para_text(p, overview)
            break


def _sow_scope(doc, content):
    scope = content.get("scope_of_work", {})
    in_s  = scope.get("in_scope", [])  if isinstance(scope, dict) else []
    out_s = scope.get("out_of_scope", []) if isinstance(scope, dict) else []
    _fill_list_group(doc, r"^In Scope Item \d+$", in_s)
    _fill_list_group(doc, r"^Out of Scope Item \d+$", out_s)


def _sow_raid(doc, content):
    def _fill_raid_tbl(tbl, items):
        if not items:
            return
        ref_tr = tbl.rows[1]._tr if len(tbl.rows) > 1 else None
        for row in list(tbl.rows)[1:]:
            tbl._tbl.remove(row._tr)
        for i, item in enumerate(items):
            if ref_tr is None:
                break
            new_tr = copy.deepcopy(ref_tr)
            tbl._tbl.append(new_tr)
            row = tbl.rows[-1]
            if isinstance(item, dict):
                _set_cell_text(row.cells[0], item.get("id", f"R-{i+1:03d}"))
                _set_cell_text(row.cells[1], item.get("description", ""))
                # Fill extra columns if the template table has them
                if len(row.cells) > 2:
                    _set_cell_text(row.cells[2], item.get("probability", item.get("impact_if_wrong", "")))
                if len(row.cells) > 3:
                    _set_cell_text(row.cells[3], item.get("impact", item.get("validation_method", "")))
                if len(row.cells) > 4:
                    _set_cell_text(row.cells[4], item.get("mitigation", item.get("resolution", "")))
            else:
                _set_cell_text(row.cells[0], f"R-{i+1:03d}")
                _set_cell_text(row.cells[1], _safe_str(item))

    _fill_raid_tbl(doc.tables[2], content.get("risks", []))
    _fill_raid_tbl(doc.tables[3], content.get("assumptions", []))
    _fill_raid_tbl(doc.tables[4], content.get("dependencies", []))


def _sow_responsibilities(doc, content):
    resp = content.get("responsibilities", {})
    invuric_items, client_items = [], []

    if isinstance(resp, dict):
        for k, v in resp.items():
            items = v if isinstance(v, list) else [str(v)]
            if any(x in k.lower() for x in ("invuric", "provider", "vendor")):
                invuric_items = [_safe_str(i) for i in items]
            elif "client" in k.lower():
                client_items = [_safe_str(i) for i in items]

    in_invuric = in_client = False
    inv_paras, cli_paras = [], []

    for p in doc.paragraphs:
        txt = p.text.strip()
        if "4.1" in txt or ("PROVIDER" in txt.upper() and "RESPONSIB" in txt.upper()):
            in_invuric, in_client = True, False
        elif "4.2" in txt or ("CLIENT" in txt.upper() and "RESPONSIB" in txt.upper()):
            in_invuric, in_client = False, True
        elif re.match(r"^(PROJECT EXECUTION|ROLES|DELIVERABLES|ENGAGEMENT)", txt.upper()):
            in_invuric = in_client = False
        elif re.match(r"^Responsibility Item \d+$", txt):
            if in_invuric:
                inv_paras.append(p)
            elif in_client:
                cli_paras.append(p)

    _replace_paras_with_items(inv_paras, invuric_items)
    _replace_paras_with_items(cli_paras, client_items)


def _sow_gantt(doc, content):
    timeline = content.get("timeline", {})
    phases   = timeline.get("phases", []) if isinstance(timeline, dict) else \
               (timeline if isinstance(timeline, list) else [])

    tbl = doc.tables[5]

    for row_idx in range(1, len(tbl.rows)):
        row = tbl.rows[row_idx]
        pi  = row_idx - 1

        if pi < len(phases):
            phase = phases[pi]
            name  = phase.get("phase", phase.get("name", f"Phase {pi+1}"))
            weeks = phase.get("weeks", phase.get("week_range", ""))
            sw, ew = _parse_week_nums(weeks, pi, len(phases))

            _set_cell_text(row.cells[0], name)
            for ci, (cs, ce) in enumerate(GANTT_COLS):
                cell = row.cells[ci + 1]
                if _weeks_overlap(sw, ew, cs, ce):
                    _set_cell_text(cell, GANTT_CHAR)
                else:
                    _set_cell_text(cell, "")
        else:
            for cell in row.cells:
                _set_cell_text(cell, "")


def _sow_roles(doc, content):
    roles = content.get("roles", [])
    tbl   = doc.tables[6]

    for row_idx in range(1, len(tbl.rows)):
        row = tbl.rows[row_idx]
        ri  = row_idx - 1

        if ri < len(roles):
            role = roles[ri]
            if isinstance(role, dict):
                _set_cell_text(row.cells[0], role.get("role", ""))
                _set_cell_text(row.cells[1], role.get("description", ""))
                _set_cell_text(row.cells[2], role.get("location", ""))
                _set_cell_text(row.cells[3], "")
                _set_cell_text(row.cells[4], "")
                _set_cell_text(row.cells[5], "")
                _set_cell_text(row.cells[6], "")
                _set_cell_text(row.cells[7], role.get("hours", ""))
                _set_cell_text(row.cells[8], "")
        else:
            for cell in row.cells:
                _set_cell_text(cell, "")


def _sow_deliverables(doc, content):
    deliverables = content.get("deliverables", [])
    acceptance   = content.get("acceptance_criteria", [])

    names = []
    if isinstance(deliverables, list) and deliverables:
        if isinstance(deliverables[0], dict):
            names = [d.get("name", d.get("description", _safe_str(d))) for d in deliverables]
        else:
            names = [_safe_str(d) for d in deliverables]

    _fill_list_group(doc, r"^Deliverable \d+$", names)

    if isinstance(acceptance, list):
        _fill_list_group(doc, r"^Acceptance Criterion \d+$", [_safe_str(a) for a in acceptance])
    elif isinstance(acceptance, str):
        _fill_list_group(doc, r"^Acceptance Criterion \d+$", [acceptance])


def _sow_signature(doc, metadata):
    # Guard against templates with fewer than 10 tables
    if len(doc.tables) <= 9:
        return
    tbl    = doc.tables[9]
    client = metadata.get("client_name", "")
    if client and len(tbl.rows) > 0 and len(tbl.rows[0].cells) > 1:
        _set_cell_text(tbl.rows[0].cells[1], client.upper())


# ═══════════════════════════════════════════════════════════════════════════════
#  Shared scratch-build helpers (PRD / FRD — match SOW template visual style)
# ═══════════════════════════════════════════════════════════════════════════════

def _setup_page(doc):
    sec = doc.sections[0]
    sec.page_width    = Inches(8.5)
    sec.page_height   = Inches(11)
    sec.left_margin   = Inches(1.0)
    sec.right_margin  = Inches(1.0)
    sec.top_margin    = Inches(1.0)
    sec.bottom_margin = Inches(1.0)
    norm = doc.styles["Normal"]
    norm.font.name = FONT
    norm.font.size = Pt(11)


def _footer(doc, doc_type: str):
    for sec in doc.sections:
        fp = sec.footer.paragraphs[0]
        fp.clear()
        run = fp.add_run(
            f"Invuric  ·  {doc_type}  ·  {datetime.utcnow().strftime('%B %Y')}  ·  CONFIDENTIAL"
        )
        _style_run(run, size=8, color_hex="888888")
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _cover(doc, title: str, metadata: dict):
    tbl = doc.add_table(rows=4, cols=1)
    tbl.style = "Table Grid"

    # Row 0: White — logo centred
    r0 = tbl.rows[0].cells[0]
    _shade(r0, HEX_WHITE)
    logo_para = r0.paragraphs[0]
    _add_invuric_logo(logo_para, width=Inches(3.0))
    logo_para.paragraph_format.space_before = Pt(24)
    logo_para.paragraph_format.space_after  = Pt(24)

    # Row 1: Navy header band — document type
    r1 = tbl.rows[1].cells[0]
    _shade(r1, HEX_NAVY)
    p1 = r1.paragraphs[0]
    p1.clear()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p1.add_run(title)
    _style_run(run, size=28, bold=True, color_hex=HEX_WHITE)
    p1.paragraph_format.space_before = Pt(20)
    p1.paragraph_format.space_after  = Pt(20)

    # Row 2: White — client name
    r2 = tbl.rows[2].cells[0]
    _shade(r2, HEX_WHITE)
    client = metadata.get("client_name", "")

    p_label = r2.paragraphs[0]
    p_label.clear()
    p_label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label_run = p_label.add_run("PREPARED FOR")
    _style_run(label_run, size=9, color_hex="94A3B8")
    p_label.paragraph_format.space_before = Pt(20)
    p_label.paragraph_format.space_after  = Pt(6)

    p_client = r2.add_paragraph()
    p_client.alignment = WD_ALIGN_PARAGRAPH.CENTER
    client_run = p_client.add_run((client or "Client Name").upper())
    _style_run(client_run, size=20, bold=True, color_hex=HEX_NAVY)
    p_client.paragraph_format.space_before = Pt(0)
    p_client.paragraph_format.space_after  = Pt(20)

    # Row 3: Light slate — prepared-by / date
    r3 = tbl.rows[3].cells[0]
    _shade(r3, HEX_SLATE)
    p3 = r3.paragraphs[0]
    p3.clear()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    prepared_run = p3.add_run("Invuric Technologies Pty. Ltd.")
    _style_run(prepared_run, size=10, bold=True, color_hex="334155")
    sep_run = p3.add_run("  ·  ")
    _style_run(sep_run, size=10, color_hex="94A3B8")
    date_run = p3.add_run(datetime.utcnow().strftime("%d %B %Y"))
    _style_run(date_run, size=10, color_hex="334155")
    p3.paragraph_format.space_before = Pt(14)
    p3.paragraph_format.space_after  = Pt(14)


def _metadata_table(doc, metadata: dict, req_label="Requestor"):
    tbl = doc.add_table(rows=4, cols=4)
    tbl.style = "Table Grid"
    _borders(tbl, "000000")

    rows_data = [
        ("Project Id",           metadata.get("project_id", ""),
         "Name of the Project",  metadata.get("project_name", "")),
        ("Author",               metadata.get("author", ""),
         req_label,              metadata.get("requestor", "")),
        ("Date Created",         datetime.utcnow().strftime("%d/%m/%Y"),
         "Date Submitted",       datetime.utcnow().strftime("%d/%m/%Y")),
        ("Estimated Start Date", metadata.get("start_date", ""),
         "Estimated End Date",   metadata.get("end_date", "")),
    ]

    for ri, (l1, v1, l2, v2) in enumerate(rows_data):
        row = tbl.rows[ri]
        _shade(row.cells[0], HEX_BLACK)
        _set_cell_text(row.cells[0], l1)
        for para in row.cells[0].paragraphs:
            for run in para.runs:
                _style_run(run, size=10, bold=True, color_hex=HEX_WHITE)

        _shade(row.cells[2], HEX_BLACK)
        _set_cell_text(row.cells[2], l2)
        for para in row.cells[2].paragraphs:
            for run in para.runs:
                _style_run(run, size=10, bold=True, color_hex=HEX_WHITE)

        _set_cell_text(row.cells[1], v1)
        _set_cell_text(row.cells[3], v2)
        for ci in (1, 3):
            for para in row.cells[ci].paragraphs:
                for run in para.runs:
                    _style_run(run, size=10)


def _section_heading(doc, text: str, number: str = ""):
    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    label = f"{number}. {text}" if number else text
    run = p.add_run(label.upper())
    _style_run(run, size=20, bold=False)
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after  = Pt(4)
    pPr  = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bot  = OxmlElement("w:bottom")
    bot.set(qn("w:val"),   "single")
    bot.set(qn("w:sz"),    "6")
    bot.set(qn("w:space"), "1")
    bot.set(qn("w:color"), HEX_BLACK)
    pBdr.append(bot)
    pPr.append(pBdr)


def _sub_heading(doc, text: str):
    p = doc.add_paragraph()
    run = p.add_run(text)
    _style_run(run, size=14, bold=True)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after  = Pt(3)


def _mini_heading(doc, text: str):
    """Small bold label for labelled content blocks."""
    p = doc.add_paragraph()
    run = p.add_run(text)
    _style_run(run, size=11, bold=True)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after  = Pt(2)


def _body(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(_safe_str(text))
    _style_run(run, size=11)
    p.paragraph_format.space_after = Pt(6)


def _bullet_list(doc, items):
    for item in (items if isinstance(items, list) else [items]):
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(_safe_str(item))
        _style_run(run, size=11)
        p.paragraph_format.space_after = Pt(3)


def _numbered_list(doc, items):
    for item in (items if isinstance(items, list) else [items]):
        p = doc.add_paragraph(style="List Number")
        run = p.add_run(_safe_str(item))
        _style_run(run, size=11)
        p.paragraph_format.space_after = Pt(3)


def _scratch_table(doc, headers, rows_data, col_widths=None):
    """Build a table from scratch matching template's black-header style."""
    if not rows_data:
        return None
    tbl = doc.add_table(rows=1 + len(rows_data), cols=len(headers))
    tbl.style = "Table Grid"
    _header_row(tbl, headers)

    for ri, row_vals in enumerate(rows_data):
        row = tbl.rows[ri + 1]
        for ci, val in enumerate(row_vals):
            if ci < len(row.cells):
                _set_cell_text(row.cells[ci], val)
                for para in row.cells[ci].paragraphs:
                    for run in para.runs:
                        _style_run(run, size=10)
        if ri % 2 == 1:
            for cell in row.cells:
                _shade(cell, HEX_LIGHT)

    _borders(tbl, "000000")
    if col_widths:
        _set_col_widths(tbl, col_widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return tbl


def _render(doc, val):
    if not val:
        return
    if isinstance(val, str):
        _body(doc, val)
    elif isinstance(val, list):
        for item in val:
            if isinstance(item, dict):
                for k, v in item.items():
                    p = doc.add_paragraph()
                    r1 = p.add_run(f"{k.replace('_',' ').title()}: ")
                    _style_run(r1, size=11, bold=True)
                    r2 = p.add_run(_safe_str(v))
                    _style_run(r2, size=11)
                    p.paragraph_format.space_after = Pt(3)
            else:
                p = doc.add_paragraph(style="List Bullet")
                run = p.add_run(_safe_str(item))
                _style_run(run, size=11)
                p.paragraph_format.space_after = Pt(3)
    elif isinstance(val, dict):
        for k, v in val.items():
            _sub_heading(doc, k.replace("_", " ").title())
            _render(doc, v)
    else:
        _body(doc, _safe_str(val))


def _signature_table(doc, metadata: dict):
    client = metadata.get("client_name", "Client")
    rows = [
        ("Name", "", "Janardan Singh"),
        ("Signature", "", ""),
        ("Date", "", ""),
    ]
    tbl = doc.add_table(rows=4, cols=3)
    tbl.style = "Table Grid"
    _header_row(tbl, ["", f"{client.upper()}", "INVURIC"])
    for ri, (lbl, cv, iv) in enumerate(rows):
        row = tbl.rows[ri + 1]
        _set_cell_text(row.cells[0], lbl)
        _set_cell_text(row.cells[1], cv)
        _set_cell_text(row.cells[2], iv)
        if ri == 1:
            row.height = Inches(0.8)
    _borders(tbl, "000000")
    _set_col_widths(tbl, [1.0, 2.75, 2.75])


# ═══════════════════════════════════════════════════════════════════════════════
#  PRD — from scratch, matching SOW template style
# ═══════════════════════════════════════════════════════════════════════════════

def _build_prd(doc, content: dict, metadata: dict):
    _cover(doc, "PRODUCT REQUIREMENTS DOCUMENT", metadata)
    doc.add_page_break()
    _metadata_table(doc, metadata, req_label="Product Owner")
    doc.add_paragraph()

    # 1. Product Overview
    _section_heading(doc, "Product Overview", "1")
    _body(doc, content.get("product_overview", ""))

    # 2. Problem Statement
    _section_heading(doc, "Problem Statement", "2")
    _body(doc, content.get("problem_statement", ""))

    # 3. Goals and Success Metrics
    _section_heading(doc, "Goals and Success Metrics", "3")
    goals = content.get("goals", [])
    if goals and isinstance(goals, list) and isinstance(goals[0], dict):
        _scratch_table(doc,
            ["Goal", "Success Metric", "Target", "Timeframe"],
            [[g.get("goal", ""), g.get("metric", ""), g.get("target", ""), g.get("timeframe", "")] for g in goals],
            [2.0, 1.8, 1.5, 1.2])
    else:
        _render(doc, goals)

    # 4. User Personas
    _section_heading(doc, "User Personas", "4")
    personas = content.get("user_personas", [])
    if personas and isinstance(personas, list):
        for i, p in enumerate(personas):
            if not isinstance(p, dict):
                _body(doc, _safe_str(p))
                continue
            _sub_heading(doc, f"Persona {i+1}: {p.get('name', '')} — {p.get('role', '')}")
            fields = [
                ("Age Range",         p.get("age_range", "")),
                ("Context",           p.get("context", "")),
                ("Goals",             p.get("goals", "")),
                ("Pain Points",       p.get("pain_points", "")),
                ("Tech Proficiency",  p.get("tech_proficiency", "")),
                ("Success Definition",p.get("success_definition", "")),
            ]
            tbl = doc.add_table(rows=len(fields), cols=2)
            tbl.style = "Table Grid"
            for ri, (label, value) in enumerate(fields):
                row = tbl.rows[ri]
                _shade(row.cells[0], HEX_LIGHT)
                _set_cell_text(row.cells[0], label)
                for para in row.cells[0].paragraphs:
                    for run in para.runs:
                        _style_run(run, size=10, bold=True)
                _set_cell_text(row.cells[1], value)
                for para in row.cells[1].paragraphs:
                    for run in para.runs:
                        _style_run(run, size=10)
            _borders(tbl, "000000")
            _set_col_widths(tbl, [1.6, 4.9])
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
    else:
        _render(doc, personas)

    # 5. Feature Requirements
    _section_heading(doc, "Feature Requirements", "5")
    features = content.get("feature_requirements", [])
    if features and isinstance(features, list) and isinstance(features[0], dict):
        for feat in features:
            _sub_heading(doc, f"{feat.get('id', '')} — {feat.get('feature', '')}")
            rows = [
                ("Priority",            feat.get("priority", "")),
                ("Description",         feat.get("description", "")),
                ("User Story",          feat.get("user_story", "")),
                ("Acceptance Criteria", feat.get("acceptance_criteria", "")),
                ("Dependencies",        feat.get("dependencies", "")),
            ]
            tbl = doc.add_table(rows=len(rows), cols=2)
            tbl.style = "Table Grid"
            for ri, (label, value) in enumerate(rows):
                row = tbl.rows[ri]
                _shade(row.cells[0], HEX_BLACK if ri == 0 else HEX_LIGHT)
                _set_cell_text(row.cells[0], label)
                for para in row.cells[0].paragraphs:
                    for run in para.runs:
                        _style_run(run, size=10, bold=True, color_hex=HEX_WHITE if ri == 0 else HEX_BLACK)
                _set_cell_text(row.cells[1], value)
                for para in row.cells[1].paragraphs:
                    for run in para.runs:
                        _style_run(run, size=10)
            _borders(tbl, "000000")
            _set_col_widths(tbl, [1.6, 4.9])
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
    else:
        _render(doc, features)

    # 6. Non-Functional Requirements
    _section_heading(doc, "Non-Functional Requirements", "6")
    nfr = content.get("non_functional_requirements", [])
    if nfr and isinstance(nfr, list) and isinstance(nfr[0], dict):
        _scratch_table(doc,
            ["Category", "Requirement", "Metric"],
            [[r.get("category", ""), r.get("requirement", ""), r.get("metric", "")] for r in nfr],
            [1.4, 2.5, 2.6])
    else:
        _render(doc, nfr)

    # 7. User Flows
    _section_heading(doc, "User Flows", "7")
    flows = content.get("user_flows", [])
    if flows and isinstance(flows, list) and isinstance(flows[0], dict):
        for flow in flows:
            _sub_heading(doc, f"{flow.get('flow_name', 'Flow')} — Actor: {flow.get('actor', '')}")
            steps = flow.get("steps", [])
            if isinstance(steps, list):
                _numbered_list(doc, steps)
            else:
                _body(doc, _safe_str(steps))
            outcome = flow.get("success_outcome", "")
            if outcome:
                p = doc.add_paragraph()
                r1 = p.add_run("Success Outcome: ")
                _style_run(r1, size=11, bold=True)
                r2 = p.add_run(_safe_str(outcome))
                _style_run(r2, size=11)
                p.paragraph_format.space_after = Pt(6)
    elif isinstance(flows, str):
        _body(doc, flows)
    else:
        _render(doc, flows)

    # 8. Edge Cases and Failure Scenarios
    _section_heading(doc, "Edge Cases and Failure Scenarios", "8")
    edge_cases = content.get("edge_cases", [])
    if edge_cases and isinstance(edge_cases, list) and isinstance(edge_cases[0], dict):
        _scratch_table(doc,
            ["Scenario", "Trigger", "Expected System Behavior", "User Communication"],
            [[e.get("scenario", ""), e.get("trigger", ""),
              e.get("expected_system_behavior", ""), e.get("user_communication", "")] for e in edge_cases],
            [1.3, 1.2, 2.2, 1.8])
    else:
        _render(doc, edge_cases)

    # 9. Risks
    _section_heading(doc, "Risks", "9")
    risks = content.get("risks", [])
    if risks and isinstance(risks, list) and isinstance(risks[0], dict):
        _scratch_table(doc,
            ["ID", "Risk", "Likelihood", "Impact", "Mitigation"],
            [[r.get("id", ""), r.get("risk", ""), r.get("likelihood", ""),
              r.get("impact", ""), r.get("mitigation", "")] for r in risks],
            [0.5, 1.8, 0.8, 0.7, 2.7])
    else:
        _render(doc, risks)

    # 10. Technical Constraints
    _section_heading(doc, "Technical Constraints", "10")
    _render(doc, content.get("technical_constraints", []))

    # 11. Out of Scope
    _section_heading(doc, "Out of Scope", "11")
    _render(doc, content.get("out_of_scope", []))

    # 12. Testing Strategy
    _section_heading(doc, "Testing Strategy", "12")
    ts = content.get("testing_strategy", {})
    if isinstance(ts, dict):
        approach = ts.get("approach", "")
        if approach:
            _body(doc, approach)
        test_types = ts.get("test_types", [])
        if test_types:
            _sub_heading(doc, "Test Types")
            _bullet_list(doc, test_types)
        coverage = ts.get("coverage_areas", [])
        if coverage:
            _sub_heading(doc, "Coverage Areas")
            _bullet_list(doc, coverage)
        vague = ts.get("vague_input_handling", "")
        if vague:
            _sub_heading(doc, "Handling Vague or Incomplete Inputs")
            _body(doc, vague)
        conflict = ts.get("conflicting_requirements_handling", "")
        if conflict:
            _sub_heading(doc, "Conflicting Requirements Handling")
            _body(doc, conflict)
    else:
        _render(doc, ts)

    # 13. Assumptions
    _section_heading(doc, "Assumptions and Constraints", "13")
    _render(doc, content.get("assumptions", []))

    # 14. Timeline
    _section_heading(doc, "Timeline", "14")
    _body(doc, content.get("timeline", ""))

    # 15. Open Questions
    _section_heading(doc, "Open Questions", "15")
    _render(doc, content.get("open_questions", []))

    # Signature
    doc.add_page_break()
    _section_heading(doc, "Approval and Sign-off", "16")
    _signature_table(doc, metadata)


# ═══════════════════════════════════════════════════════════════════════════════
#  FRD — from scratch, matching SOW template style
# ═══════════════════════════════════════════════════════════════════════════════

def _build_frd(doc, content: dict, metadata: dict):
    _cover(doc, "FUNCTIONAL REQUIREMENTS DOCUMENT", metadata)
    doc.add_page_break()
    _metadata_table(doc, metadata, req_label="FRD Requestor")
    doc.add_paragraph()

    # 1. Introduction
    _section_heading(doc, "Introduction and Purpose", "1")
    _render(doc, content.get("introduction", content.get("introduction_and_purpose", "")))

    # 2. System Overview
    _section_heading(doc, "System Overview", "2")
    _render(doc, content.get("system_overview", ""))

    # 3. Use Cases
    _section_heading(doc, "Use Cases", "3")
    use_cases = content.get("use_cases", [])
    if use_cases and isinstance(use_cases, list):
        for uc in use_cases:
            if not isinstance(uc, dict):
                _body(doc, _safe_str(uc))
                continue
            _sub_heading(doc, f"{uc.get('id', '')} — {uc.get('name', '')}")
            p = doc.add_paragraph()
            r1 = p.add_run("Actor: ")
            _style_run(r1, size=11, bold=True)
            r2 = p.add_run(_safe_str(uc.get("actor", "")))
            _style_run(r2, size=11)

            precond = uc.get("preconditions", [])
            if precond:
                _mini_heading(doc, "Preconditions")
                _bullet_list(doc, precond if isinstance(precond, list) else [precond])

            main_flow = uc.get("main_flow", [])
            if main_flow:
                _mini_heading(doc, "Main Flow")
                _numbered_list(doc, main_flow if isinstance(main_flow, list) else [main_flow])

            alt_flows = uc.get("alternate_flows", [])
            if alt_flows:
                _mini_heading(doc, "Alternate Flows / Exception Paths")
                _bullet_list(doc, alt_flows if isinstance(alt_flows, list) else [alt_flows])

            postcond = uc.get("postconditions", [])
            if postcond:
                _mini_heading(doc, "Postconditions")
                _bullet_list(doc, postcond if isinstance(postcond, list) else [postcond])

            doc.add_paragraph().paragraph_format.space_after = Pt(4)
    else:
        _render(doc, use_cases)

    # 4. Functional Requirements
    _section_heading(doc, "Functional Requirements", "4")
    frs = content.get("functional_requirements", [])
    if frs and isinstance(frs, list) and isinstance(frs[0], dict):
        for fr in frs:
            _sub_heading(doc, f"{fr.get('id', '')} — {fr.get('title', '')}")
            rows = [
                ("Priority",            fr.get("priority", "")),
                ("Description",         fr.get("description", "")),
                ("Input",               fr.get("input", "")),
                ("Process",             fr.get("process", "")),
                ("Output",              fr.get("output", "")),
                ("Acceptance Criteria", fr.get("acceptance_criteria", "")),
            ]
            tbl = doc.add_table(rows=len(rows), cols=2)
            tbl.style = "Table Grid"
            for ri, (label, value) in enumerate(rows):
                row = tbl.rows[ri]
                _shade(row.cells[0], HEX_BLACK if ri == 0 else HEX_LIGHT)
                _set_cell_text(row.cells[0], label)
                for para in row.cells[0].paragraphs:
                    for run in para.runs:
                        _style_run(run, size=10, bold=True,
                                   color_hex=HEX_WHITE if ri == 0 else HEX_BLACK)
                _set_cell_text(row.cells[1], value)
                for para in row.cells[1].paragraphs:
                    for run in para.runs:
                        _style_run(run, size=10)
            _borders(tbl, "000000")
            _set_col_widths(tbl, [1.6, 4.9])
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
    else:
        _render(doc, frs)

    # 5. Business Rules
    _section_heading(doc, "Business Rules", "5")
    biz = content.get("business_rules", [])
    if biz and isinstance(biz, list) and isinstance(biz[0], dict):
        _scratch_table(doc,
            ["ID", "Category", "Rule", "Rationale"],
            [[b.get("id", ""), b.get("category", ""), b.get("rule", ""), b.get("rationale", "")] for b in biz],
            [0.6, 1.0, 3.0, 1.9])
    elif biz and isinstance(biz, list):
        _scratch_table(doc,
            ["ID", "Business Rule"],
            [[f"BR-{i+1:03d}", _safe_str(r)] for i, r in enumerate(biz)],
            [0.8, 5.7])
    else:
        _render(doc, biz)

    # 6. Error Handling
    _section_heading(doc, "Error Handling", "6")
    errors = content.get("error_handling", [])
    if errors and isinstance(errors, list) and isinstance(errors[0], dict):
        _scratch_table(doc,
            ["Error Code", "Scenario", "Trigger Condition", "System Response", "User Message", "Recovery Action"],
            [[e.get("error_code", ""), e.get("scenario", ""), e.get("trigger_condition", ""),
              e.get("system_response", ""), e.get("user_message", ""), e.get("recovery_action", "")] for e in errors],
            [0.75, 0.95, 1.1, 1.1, 1.2, 1.4])
    else:
        _render(doc, errors)

    # 7. Data Requirements
    _section_heading(doc, "Data Requirements", "7")
    data_req = content.get("data_requirements", [])
    if data_req and isinstance(data_req, list) and isinstance(data_req[0], dict):
        _scratch_table(doc,
            ["Entity", "Field", "Data Type", "Constraints", "Validation Rules", "Source"],
            [[d.get("entity", ""), d.get("field", ""), d.get("data_type", ""),
              d.get("constraints", ""), d.get("validation_rules", ""), d.get("source", "")] for d in data_req],
            [0.9, 0.9, 1.0, 1.0, 1.5, 0.8])
    else:
        _render(doc, content.get("data_requirements", ""))

    # 8. Integration Requirements
    _section_heading(doc, "Integration Requirements", "8")
    integ = content.get("integration_requirements", [])
    if integ and isinstance(integ, list) and isinstance(integ[0], dict):
        _scratch_table(doc,
            ["System", "Type", "Description", "Protocol", "Authentication", "Error Handling"],
            [[i.get("system", ""), i.get("type", ""), i.get("description", ""),
              i.get("protocol", ""), i.get("authentication", ""), i.get("error_handling", "")] for i in integ],
            [0.9, 0.7, 1.3, 0.8, 1.0, 1.8])
    else:
        _render(doc, integ)

    # 9-11
    for num, title, key in [
        ("9",  "Security Requirements",   "security_requirements"),
        ("10", "Reporting Requirements",  "reporting_requirements"),
        ("11", "Assumptions",             "assumptions"),
    ]:
        _section_heading(doc, title, num)
        _render(doc, content.get(key, ""))

    # 12. Glossary
    _section_heading(doc, "Glossary", "12")
    glossary = content.get("glossary", {})
    if isinstance(glossary, dict) and glossary:
        rows = [[term, defn] for term, defn in glossary.items()]
        _scratch_table(doc, ["Term", "Definition"], rows, [1.5, 5.0])
    else:
        _render(doc, glossary)

    # Signature
    doc.add_page_break()
    _section_heading(doc, "Approval and Sign-off", "13")
    _signature_table(doc, metadata)


# ═══════════════════════════════════════════════════════════════════════════════
#  Public entry point
# ═══════════════════════════════════════════════════════════════════════════════

def _flatten_template_value(value, level: int = 0) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, list):
        flattened = [_flatten_template_value(item, level + 1) for item in value]
        return "\n".join(item for item in flattened if item)
    if isinstance(value, dict):
        parts = []
        for key, item in value.items():
            rendered = _flatten_template_value(item, level + 1)
            if rendered:
                parts.append(f"{key.replace('_', ' ').title()}: {rendered}")
        return "\n".join(parts)
    return str(value)


def _iter_template_paragraphs(doc):
    for paragraph in doc.paragraphs:
        yield paragraph
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph
    for section in doc.sections:
        for paragraph in section.header.paragraphs:
            yield paragraph
        for paragraph in section.footer.paragraphs:
            yield paragraph


def _pick_style_name(doc, preferred_names: list[str]) -> str | None:
    style_names = {style.name for style in doc.styles}
    for name in preferred_names:
        if name in style_names:
            return name
    return None


def _build_template_replacements(content: dict) -> dict[str, str]:
    metadata = content.get("metadata", {})
    replacements: dict[str, str] = {}
    alias_map = {
        "client_name": ["client_name", "client", "customer_name", "customer"],
        "project_name": ["project_name", "project", "document_name"],
        "project_id": ["project_id", "project_code"],
        "author": ["author", "prepared_by"],
        "requestor": ["requestor", "requested_by", "product_owner"],
        "start_date": ["start_date", "estimated_start_date"],
        "end_date": ["end_date", "estimated_end_date"],
    }

    for key, aliases in alias_map.items():
        value = _safe_str(metadata.get(key, ""))
        for alias in aliases:
            replacements[alias] = value

    for key, value in content.items():
        if key == "metadata":
            continue
        replacements[key] = _flatten_template_value(value)

    return replacements


def _apply_template_replacements(doc, replacements: dict[str, str]):
    patterns = {}
    for token, value in replacements.items():
        patterns[f"{{{{{token}}}}}"] = value
        patterns[f"[[{token}]]"] = value
        patterns[f"[{token}]"] = value

    for paragraph in _iter_template_paragraphs(doc):
        text = paragraph.text
        updated = text
        for token, value in patterns.items():
            if token in updated:
                updated = updated.replace(token, value)
        if updated != text:
            _set_para_text(paragraph, updated)


def _client_heading(doc, text: str, style_name: str | None):
    paragraph = doc.add_paragraph()
    if style_name:
        paragraph.style = style_name
    run = paragraph.add_run(text)
    if not style_name:
        _style_run(run, size=16, bold=True)
        paragraph.paragraph_format.space_before = Pt(12)
        paragraph.paragraph_format.space_after = Pt(6)


def _client_body(doc, text: str, style_name: str | None):
    paragraph = doc.add_paragraph()
    if style_name:
        paragraph.style = style_name
    run = paragraph.add_run(_safe_str(text))
    if not style_name:
        _style_run(run, size=11)
        paragraph.paragraph_format.space_after = Pt(6)


def _client_list(doc, items, style_name: str | None):
    for item in (items if isinstance(items, list) else [items]):
        paragraph = doc.add_paragraph(style=style_name or None)
        run = paragraph.add_run(f"• {_safe_str(item)}")
        if not style_name:
            _style_run(run, size=11)
            paragraph.paragraph_format.space_after = Pt(3)


def _client_render(doc, value, heading_style: str | None, body_style: str | None):
    if not value:
        return
    if isinstance(value, str):
        _client_body(doc, value, body_style)
        return
    if isinstance(value, list):
        if value and isinstance(value[0], dict):
            for index, item in enumerate(value, start=1):
                label = item.get("name") or item.get("title") or item.get("id") or f"Item {index}"
                _client_heading(doc, _safe_str(label), heading_style)
                _client_render(doc, {k: v for k, v in item.items() if k not in {"name", "title", "id"}}, None, body_style)
        else:
            _client_list(doc, value, body_style)
        return
    if isinstance(value, dict):
        for key, item in value.items():
            _client_heading(doc, key.replace("_", " ").title(), heading_style)
            _client_render(doc, item, None, body_style)
        return
    _client_body(doc, str(value), body_style)


def _client_sections(doc_type: str, content: dict) -> list[tuple[str, object]]:
    if doc_type == "SOW":
        return [
            ("Executive Summary", content.get("executive_summary")),
            ("Project Overview", content.get("project_overview")),
            ("Scope Of Work", content.get("scope_of_work")),
            ("Deliverables", content.get("deliverables")),
            ("Risks", content.get("risks")),
            ("Assumptions", content.get("assumptions")),
            ("Dependencies", content.get("dependencies")),
            ("Responsibilities", content.get("responsibilities")),
            ("Timeline", content.get("timeline")),
            ("Roles", content.get("roles")),
            ("Change Management", content.get("change_management")),
            ("Acceptance Criteria", content.get("acceptance_criteria")),
            ("Validity", content.get("validity")),
        ]
    if doc_type == "PRD":
        return [
            ("Product Overview", content.get("product_overview")),
            ("Problem Statement", content.get("problem_statement")),
            ("Goals", content.get("goals")),
            ("User Personas", content.get("user_personas")),
            ("Feature Requirements", content.get("feature_requirements")),
            ("Non Functional Requirements", content.get("non_functional_requirements")),
            ("User Flows", content.get("user_flows")),
            ("Edge Cases", content.get("edge_cases")),
            ("Risks", content.get("risks")),
            ("Technical Constraints", content.get("technical_constraints")),
            ("Out Of Scope", content.get("out_of_scope")),
            ("Testing Strategy", content.get("testing_strategy")),
            ("Assumptions", content.get("assumptions")),
            ("Timeline", content.get("timeline")),
            ("Open Questions", content.get("open_questions")),
        ]
    return [
        ("Introduction", content.get("introduction") or content.get("introduction_and_purpose")),
        ("System Overview", content.get("system_overview")),
        ("Use Cases", content.get("use_cases")),
        ("Functional Requirements", content.get("functional_requirements")),
        ("Business Rules", content.get("business_rules")),
        ("Error Handling", content.get("error_handling")),
        ("Data Requirements", content.get("data_requirements")),
        ("Integration Requirements", content.get("integration_requirements")),
        ("Security Requirements", content.get("security_requirements")),
        ("Reporting Requirements", content.get("reporting_requirements")),
        ("Assumptions", content.get("assumptions")),
        ("Glossary", content.get("glossary")),
    ]


def _build_client_template_docx(content: dict, output_path: str, doc_type: str, template_bytes: bytes | None, template_kind: str | None):
    if template_kind == "docx" and template_bytes:
        doc = Document(io.BytesIO(template_bytes))
    else:
        doc = Document()
        _setup_page(doc)

    heading_style = _pick_style_name(doc, ["Heading 1", "Title"])
    body_style = _pick_style_name(doc, ["Body Text", "Normal"])

    _apply_template_replacements(doc, _build_template_replacements(content))
    doc.add_page_break()
    _client_heading(doc, f"{doc_type} Generated Content", heading_style)

    for title, value in _client_sections(doc_type, content):
        if not value:
            continue
        _client_heading(doc, title, heading_style)
        _client_render(doc, value, None, body_style)

    doc.save(output_path)


def generate_docx(
    content: dict,
    output_path: str,
    doc_type: str = "SOW",
    template_mode: str = "invuric",
    template_bytes: bytes | None = None,
    template_kind: str | None = None,
):
    if template_mode == "client":
        _build_client_template_docx(content, output_path, doc_type, template_bytes, template_kind)
        return

    if doc_type == "SOW":
        _sow_from_template(content, output_path)
    else:
        doc = Document()
        _setup_page(doc)
        metadata = content.get("metadata", {})
        if doc_type == "PRD":
            _build_prd(doc, content, metadata)
        elif doc_type == "FRD":
            _build_frd(doc, content, metadata)
        _footer(doc, doc_type)
        doc.save(output_path)
