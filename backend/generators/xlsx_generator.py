from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ── Shared helpers ────────────────────────────────────────────────────────────

DARK_FILL   = PatternFill("solid", fgColor="1a1a2e")
ACCENT_FILL = PatternFill("solid", fgColor="6366f1")
LIGHT_FILL  = PatternFill("solid", fgColor="f4f4f8")
ALT_FILL    = PatternFill("solid", fgColor="eef0fb")

WHITE_FONT  = Font(name="Calibri", color="FFFFFF", bold=True, size=11)
HEADER_FONT = Font(name="Calibri", bold=True, size=10)
BODY_FONT   = Font(name="Calibri", size=10)
THIN_BORDER = Border(
    left=Side(style="thin", color="d1d5db"),
    right=Side(style="thin", color="d1d5db"),
    top=Side(style="thin", color="d1d5db"),
    bottom=Side(style="thin", color="d1d5db"),
)

PRIORITY_COLORS = {
    "Critical":   "ef4444",
    "High":       "f97316",
    "Medium":     "f59e0b",
    "Low":        "22c55e",
    "Must Have":  "ef4444",
    "Should Have":"f97316",
    "Could Have": "f59e0b",
    "Won't Have": "94a3b8",
    "Internal":   "6366f1",
    "External":   "f59e0b",
    "On Track":   "22c55e",
    "At Risk":    "f59e0b",
    "Blocked":    "ef4444",
}


def _safe_str(val) -> str:
    if val is None:
        return ""
    if isinstance(val, list):
        return "\n".join(str(v) for v in val if v)
    return str(val)


def _write_header_row(ws, columns: list[str], row: int = 1):
    for col, label in enumerate(columns, 1):
        cell = ws.cell(row=row, column=col, value=label)
        cell.font = WHITE_FONT
        cell.fill = ACCENT_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
    ws.row_dimensions[row].height = 26


def _write_cell(ws, row: int, col: int, value, fill_hex: str = None, bold: bool = False):
    cell = ws.cell(row=row, column=col, value=_safe_str(value))
    cell.font = Font(name="Calibri", size=10, bold=bold,
                     color="FFFFFF" if fill_hex and int(fill_hex, 16) < 0x888888 else "1a1a2e")
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    cell.border = THIN_BORDER
    if fill_hex:
        cell.fill = PatternFill("solid", fgColor=fill_hex)
    return cell


def _write_data_row(ws, row: int, values: list, alt: bool = False):
    """Write a full data row with optional alternating fill."""
    for col, val in enumerate(values, 1):
        cell = ws.cell(row=row, column=col, value=_safe_str(val))
        cell.font = BODY_FONT
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.border = THIN_BORDER
        if alt:
            cell.fill = ALT_FILL
    return ws.row_dimensions[row]


def _color_cell(ws, row: int, col: int, value):
    """Write a cell and apply a color fill based on the value (e.g. High/Medium/Low)."""
    fill_hex = PRIORITY_COLORS.get(_safe_str(value))
    cell = ws.cell(row=row, column=col, value=_safe_str(value))
    cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF" if fill_hex else "1a1a2e")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
    cell.border = THIN_BORDER
    if fill_hex:
        cell.fill = PatternFill("solid", fgColor=fill_hex)
    return cell


def _auto_width(ws, min_w: int = 12, max_w: int = 45):
    """Auto-size columns, safely handling None cell values."""
    for col in ws.columns:
        length = 0
        for cell in col:
            try:
                val = cell.value
                cell_len = len(str(val)) if val is not None else 0
                # For wrapped text, only count the longest line
                if "\n" in str(val or ""):
                    cell_len = max(len(line) for line in str(val).split("\n"))
                length = max(length, cell_len)
            except Exception:
                pass
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(length + 2, min_w), max_w)


def _freeze_header(ws):
    ws.freeze_panes = "A2"


# ── RAID ──────────────────────────────────────────────────────────────────────

def generate_raid_xlsx(data: dict, output_path: str):
    wb = Workbook()
    wb.remove(wb.active)

    # ── Risks tab ────────────────────────────────────────────────────────────
    ws_risks = wb.create_sheet("Risks")
    risk_cols = ["ID", "Title", "Description", "Probability", "Impact", "Risk Score",
                 "Trigger Conditions", "Mitigation", "Contingency Plan", "Owner", "Review Date"]
    _write_header_row(ws_risks, risk_cols)
    for r, item in enumerate(data.get("risks", []), 2):
        alt = r % 2 == 0
        ws_risks.cell(row=r, column=1, value=item.get("id", "")).font = BODY_FONT
        ws_risks.cell(row=r, column=1).border = THIN_BORDER
        ws_risks.cell(row=r, column=1).alignment = Alignment(vertical="top")

        _write_data_row(ws_risks, r, [
            item.get("id", ""),
            item.get("title", ""),
            item.get("description", ""),
            "", "", "",  # probability, impact, risk_score — colored separately
            item.get("trigger_conditions", ""),
            item.get("mitigation", ""),
            item.get("contingency", ""),
            item.get("owner", ""),
            item.get("review_date", ""),
        ], alt=alt)
        _color_cell(ws_risks, r, 4, item.get("probability", ""))
        _color_cell(ws_risks, r, 5, item.get("impact", ""))
        _color_cell(ws_risks, r, 6, item.get("risk_score", ""))
        ws_risks.row_dimensions[r].height = 55
    _auto_width(ws_risks)
    _freeze_header(ws_risks)

    # ── Assumptions tab ──────────────────────────────────────────────────────
    ws_assum = wb.create_sheet("Assumptions")
    assum_cols = ["ID", "Title", "Description", "Impact If Wrong",
                  "Validation Method", "Validate By", "Owner"]
    _write_header_row(ws_assum, assum_cols)
    for r, item in enumerate(data.get("assumptions", []), 2):
        _write_data_row(ws_assum, r, [
            item.get("id", ""),
            item.get("title", ""),
            item.get("description", ""),
            item.get("impact_if_wrong", ""),
            item.get("validation_method", ""),
            item.get("validation_by", ""),
            item.get("owner", ""),
        ], alt=r % 2 == 0)
        ws_assum.row_dimensions[r].height = 55
    _auto_width(ws_assum)
    _freeze_header(ws_assum)

    # ── Issues tab ───────────────────────────────────────────────────────────
    ws_issues = wb.create_sheet("Issues")
    issue_cols = ["ID", "Title", "Description", "Severity", "Impact",
                  "Resolution Plan", "Resolution Owner", "Target Resolution Date"]
    _write_header_row(ws_issues, issue_cols)
    for r, item in enumerate(data.get("issues", []), 2):
        _write_data_row(ws_issues, r, [
            item.get("id", ""),
            item.get("title", ""),
            item.get("description", ""),
            "",  # severity — colored separately
            item.get("impact", ""),
            item.get("resolution_plan", item.get("resolution", "")),
            item.get("resolution_owner", item.get("owner", "")),
            item.get("target_resolution_date", ""),
        ], alt=r % 2 == 0)
        _color_cell(ws_issues, r, 4, item.get("severity", ""))
        ws_issues.row_dimensions[r].height = 55
    _auto_width(ws_issues)
    _freeze_header(ws_issues)

    # ── Dependencies tab ─────────────────────────────────────────────────────
    ws_deps = wb.create_sheet("Dependencies")
    dep_cols = ["ID", "Title", "Description", "Type", "Due Date",
                "Owner", "Impact If Delayed", "Status"]
    _write_header_row(ws_deps, dep_cols)
    for r, item in enumerate(data.get("dependencies", []), 2):
        _write_data_row(ws_deps, r, [
            item.get("id", ""),
            item.get("title", ""),
            item.get("description", ""),
            "",  # type — colored
            item.get("due_date", ""),
            item.get("dependency_owner", item.get("owner", "")),
            item.get("impact_if_delayed", ""),
            "",  # status — colored
        ], alt=r % 2 == 0)
        _color_cell(ws_deps, r, 4, item.get("type", ""))
        _color_cell(ws_deps, r, 8, item.get("status", ""))
        ws_deps.row_dimensions[r].height = 55
    _auto_width(ws_deps)
    _freeze_header(ws_deps)

    wb.save(output_path)


# ── WBS ───────────────────────────────────────────────────────────────────────

def generate_wbs_xlsx(data: dict, output_path: str, audience: list[str]):
    wb = Workbook()
    ws = wb.active
    ws.title = "WBS"

    columns = ["ID", "Phase / Task / Subtask", "Description", "Duration", "Assigned To", "Dependencies", "Level"]
    _write_header_row(ws, columns)

    PHASE_FILL   = PatternFill("solid", fgColor="1a1a2e")
    TASK_FILL    = PatternFill("solid", fgColor="dde0f7")
    SUBTASK_FILL = PatternFill("solid", fgColor="f4f4f8")

    row = 2
    for phase in data.get("phases", []):
        # Phase row
        phase_vals = [
            phase.get("id", ""),
            phase.get("name", ""),
            f"Objective: {phase.get('objective', '')} | Exit: {phase.get('exit_criteria', '')}",
            phase.get("duration", ""),
            "",
            "",
            "Phase",
        ]
        for c, val in enumerate(phase_vals, 1):
            cell = ws.cell(row=row, column=c, value=val)
            cell.font = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
            cell.fill = PHASE_FILL
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 30
        row += 1

        for task in phase.get("tasks", []):
            assigned = ", ".join(a for a in task.get("assigned_to", []) if not audience or a in audience)
            task_vals = [
                task.get("id", ""),
                f"  {task.get('name', '')}",
                task.get("description", ""),
                task.get("duration", ""),
                assigned,
                task.get("dependencies", ""),
                "Task",
            ]
            for c, val in enumerate(task_vals, 1):
                cell = ws.cell(row=row, column=c, value=val)
                cell.font = HEADER_FONT
                cell.fill = TASK_FILL
                cell.border = THIN_BORDER
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            ws.row_dimensions[row].height = 40
            row += 1

            for subtask in task.get("subtasks", []):
                s_assigned = ", ".join(a for a in subtask.get("assigned_to", []) if not audience or a in audience)
                sub_vals = [
                    subtask.get("id", ""),
                    f"    {subtask.get('name', '')}",
                    "",
                    subtask.get("duration", ""),
                    s_assigned,
                    "",
                    "Subtask",
                ]
                for c, val in enumerate(sub_vals, 1):
                    cell = ws.cell(row=row, column=c, value=val)
                    cell.font = BODY_FONT
                    cell.fill = SUBTASK_FILL
                    cell.border = THIN_BORDER
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                ws.row_dimensions[row].height = 25
                row += 1

    # Summary row at bottom
    ws.cell(row=row, column=1, value="SUMMARY").font = Font(name="Calibri", bold=True, size=10, color="FFFFFF")
    ws.cell(row=row, column=1).fill = ACCENT_FILL
    ws.cell(row=row, column=2, value=f"Total: {data.get('total_phases',0)} Phases | {data.get('total_tasks',0)} Tasks | {data.get('total_subtasks',0)} Subtasks").font = HEADER_FONT
    for c in range(1, len(columns) + 1):
        ws.cell(row=row, column=c).border = THIN_BORDER

    _auto_width(ws)
    _freeze_header(ws)
    wb.save(output_path)


# ── Backlog / User Stories ────────────────────────────────────────────────────

BACKLOG_COLS = [
    "Page Name", "Epic", "High Level Flow", "User Story",
    "Priority", "Story Points",
    "Acceptance Criteria", "Data Points", "Edge Cases",
    "Non-Functional Requirements", "Dependencies",
    "Project ID", "Project Name", "Analyzed At",
]
BACKLOG_FIELDS = [
    "page_name", "epic", "high_level_flow", "user_story",
    "priority", "story_points",
    "acceptance_criteria", "data_points", "edge_cases",
    "non_functional", "dependencies",
    "project_id", "project_name", "analyzed_at",
]


def generate_backlog_xlsx(data: dict, output_path: str, project_name: str, project_id: str):
    wb = Workbook()
    ws = wb.active
    ws.title = "User Stories"

    _write_header_row(ws, BACKLOG_COLS)

    for r, story in enumerate(data.get("stories", []), 2):
        alt = r % 2 == 0
        for c, field in enumerate(BACKLOG_FIELDS, 1):
            value = story.get(field, "")
            cell = ws.cell(row=r, column=c)

            # Priority and Story Points get special formatting
            if field == "priority":
                _color_cell(ws, r, c, value)
            elif field == "story_points":
                cell.value = value
                cell.font = Font(name="Calibri", bold=True, size=11)
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = THIN_BORDER
                if alt:
                    cell.fill = ALT_FILL
            else:
                cell.value = _safe_str(value)
                cell.font = BODY_FONT
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = THIN_BORDER
                if alt:
                    cell.fill = ALT_FILL

        ws.row_dimensions[r].height = 80

    # Set column widths manually for readability (backlog has many cols)
    manual_widths = {
        1: 18,   # Page Name
        2: 20,   # Epic
        3: 28,   # High Level Flow
        4: 45,   # User Story
        5: 14,   # Priority
        6: 8,    # Story Points
        7: 45,   # Acceptance Criteria
        8: 30,   # Data Points
        9: 35,   # Edge Cases
        10: 30,  # Non-Functional
        11: 20,  # Dependencies
        12: 14,  # Project ID
        13: 22,  # Project Name
        14: 18,  # Analyzed At
    }
    for col_idx, width in manual_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    _freeze_header(ws)
    wb.save(output_path)
